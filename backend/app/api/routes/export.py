from fastapi import APIRouter, Depends, HTTPException
from pydantic import BaseModel
from sqlalchemy.orm import Session
from sqlalchemy import text
from pathlib import Path
from fastapi.responses import FileResponse

from app.db.session import get_db
from app import models
from app.core.security import get_current_user

router = APIRouter()

class ExportRequest(BaseModel):
    format: str = "pdf"
    citation_style: str = "APA"  # or IEEE


def ensure_export_dir() -> Path:
    p = Path('/tmp/tesispilot_exports')
    p.mkdir(parents=True, exist_ok=True)
    return p


def format_references(refs, style: str = "APA") -> str:
    lines = []
    for r in refs:
        # r may be a Row or tuple; try dict-like
        try:
            title = r['title']
            authors = r['authors'] or ''
            year = r['year'] or ''
            source = r['source'] or ''
            doi = r['doi'] or ''
            url = r['url'] or ''
        except Exception:
            # fallback positional
            title, authors, year, source, doi, url = (r[0], r[1], r[2], r[3] if len(r) > 3 else '', r[4] if len(r) > 4 else '', r[5] if len(r) > 5 else '')

        if style.upper() == 'IEEE':
            parts = []
            if authors:
                parts.append(authors)
            if title:
                parts.append(f'"{title}"')
            if source:
                parts.append(source)
            if year:
                parts.append(str(year))
            if doi:
                parts.append(f'DOI: {doi}')
            elif url:
                parts.append(url)
            lines.append(', '.join(parts) + '.')
        else:  # default APA-like
            # APA: Authors (Year). Title. Source. DOI/URL
            author_part = f"{authors}." if authors else ''
            year_part = f"({year})." if year else ''
            title_part = f"{title}." if title else ''
            source_part = f"{source}." if source else ''
            link_part = ''
            if doi:
                link_part = f"https://doi.org/{doi}"
            elif url:
                link_part = url
            if link_part:
                lines.append(f"{author_part} {year_part} {title_part} {source_part} {link_part}")
            else:
                lines.append(f"{author_part} {year_part} {title_part} {source_part}")
    return '\n'.join(lines)


@router.post('/theses/{thesis_id}/export')
def export_thesis(thesis_id: str, payload: ExportRequest, db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):
    thesis = db.query(models.Thesis).filter(models.Thesis.id == thesis_id).first()
    if not thesis:
        raise HTTPException(status_code=404, detail='Thesis not found')
    if thesis.user_id != current_user.id:
        raise HTTPException(status_code=403, detail='Forbidden')

    chapters = db.query(models.Chapter).filter(models.Chapter.thesis_id == thesis_id).order_by(models.Chapter.order_index).all()

    # start building markdown content
    content = f"# {thesis.title}\n\n"
    content += f"## Resumen\n{thesis.abstract or 'Sin resumen.'}\n\n"
    content += f"## Problema de investigación\n{thesis.problem_statement or 'Sin problema de investigación definido.'}\n\n"
    content += f"## Objetivos\n{thesis.objectives or 'Sin objetivos definidos.'}\n\n"

    for ch in chapters:
        content += f"## {ch.title}\n{ch.content or ''}\n\n"

    # fetch references if table exists
    refs = []
    try:
        q = text("SELECT title, authors, year, source, doi, url FROM references WHERE thesis_id = :tid")
        res = db.execute(q, {"tid": thesis_id})
        refs = res.fetchall()
    except Exception:
        refs = []

    if refs:
        formatted = format_references(refs, payload.citation_style or 'APA')
        content += f"## Referencias\n{formatted}\n\n"

    # fetch annexes/document files
    annexes = []
    try:
        q2 = text("SELECT file_name, file_url FROM document_files WHERE thesis_id = :tid")
        r2 = db.execute(q2, {"tid": thesis_id})
        annexes = r2.fetchall()
    except Exception:
        annexes = []

    if annexes:
        content += "## Anexos\n"
        for a in annexes:
            try:
                fname = a['file_name']
                furl = a['file_url']
            except Exception:
                fname, furl = (a[0], a[1] if len(a) > 1 else '')
            content += f"- {fname} ({furl})\n"
        content += "\n"

    export_dir = ensure_export_dir()
    safe_name = ''.join(c if c.isalnum() or c in ('-', '_') else '_' for c in thesis.title.lower())[:80]
    out_path = export_dir / f"{safe_name}.{payload.format}"

    if payload.format == 'md':
        out_path.write_text(content, encoding='utf-8')
    elif payload.format == 'docx':
        try:
            from docx import Document
            doc = Document()
            doc.add_heading(thesis.title, level=1)
            doc.add_paragraph(thesis.abstract or 'Sin resumen.')
            doc.add_paragraph('\n')
            doc.add_heading('Problema de investigación', level=2)
            doc.add_paragraph(thesis.problem_statement or '')
            doc.add_heading('Objetivos', level=2)
            doc.add_paragraph(thesis.objectives or '')
            for ch in chapters:
                doc.add_heading(ch.title, level=2)
                doc.add_paragraph(ch.content or '')
            if refs:
                doc.add_heading('Referencias', level=2)
                formatted = format_references(refs, payload.citation_style or 'APA')
                for line in formatted.split('\n'):
                    doc.add_paragraph(line)
            if annexes:
                doc.add_heading('Anexos', level=2)
                for a in annexes:
                    try:
                        fname = a['file_name']
                        furl = a['file_url']
                    except Exception:
                        fname, furl = (a[0], a[1] if len(a) > 1 else '')
                    doc.add_paragraph(f"{fname} ({furl})")
            doc.save(str(out_path))
        except Exception as exc:
            raise HTTPException(status_code=500, detail=f'Could not generate DOCX: {exc}')
    elif payload.format == 'pdf':
        try:
            from reportlab.lib.pagesizes import letter
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.lib.units import inch
            styles = getSampleStyleSheet()
            doc = SimpleDocTemplate(str(out_path), pagesize=letter)
            story = []
            story.append(Paragraph(thesis.title, styles['Title']))
            story.append(Spacer(1, 0.25 * inch))
            story.append(Paragraph(f"Resumen: {thesis.abstract or 'Sin resumen.'}", styles['BodyText']))
            story.append(Spacer(1, 0.20 * inch))
            story.append(Paragraph('Problema de investigación', styles['Heading2']))
            story.append(Paragraph(thesis.problem_statement or '', styles['BodyText']))
            story.append(Spacer(1, 0.15 * inch))
            story.append(Paragraph('Objetivos', styles['Heading2']))
            story.append(Paragraph(thesis.objectives or '', styles['BodyText']))
            story.append(Spacer(1, 0.15 * inch))
            for ch in chapters:
                story.append(Paragraph(ch.title, styles['Heading2']))
                story.append(Paragraph(ch.content or '', styles['BodyText']))
                story.append(Spacer(1, 0.15 * inch))
            if refs:
                story.append(Paragraph('Referencias', styles['Heading2']))
                formatted = format_references(refs, payload.citation_style or 'APA')
                for line in formatted.split('\n'):
                    story.append(Paragraph(line, styles['BodyText']))
            if annexes:
                story.append(Paragraph('Anexos', styles['Heading2']))
                for a in annexes:
                    try:
                        fname = a['file_name']
                        furl = a['file_url']
                    except Exception:
                        fname, furl = (a[0], a[1] if len(a) > 1 else '')
                    story.append(Paragraph(f"{fname} ({furl})", styles['BodyText']))
            doc.build(story)
        except Exception as exc:
            raise HTTPException(status_code=500, detail=f'Could not generate PDF: {exc}')
    else:
        raise HTTPException(status_code=400, detail='Unsupported format. Use md, docx, or pdf')

    return {'filename': out_path.name, 'download_url': f'/api/download/{out_path.name}'}


@router.get('/download/{filename}')
def download_file(filename: str):
    file_path = Path('/tmp/tesispilot_exports') / filename
    if not file_path.exists():
        raise HTTPException(status_code=404, detail='File not found')
    return FileResponse(file_path, filename=filename)
